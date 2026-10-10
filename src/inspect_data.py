
from pathlib import Path
import sqlite3

DATA_DIR = Path(__file__).resolve().parent.parent / "data"


def inspect_files():
    print("\n=== RIDEWISE DATA INSPECTION ===")

    print("\n1. Files found in data/")

    for file_path in sorted(DATA_DIR.iterdir()):
        if file_path.is_file():
            print(f"  {file_path.name}")

    # Check whether the required SQLite database exists
    db_path = DATA_DIR / "ridewise.db"

    print("\n2. SQLite database")

    if not db_path.exists():
        print("  ridewise.db was not found.")
        print("  We will need to locate the supplied database.")
    else:
        connection = sqlite3.connect(
            f"file:{db_path.as_posix()}?mode=ro",
            uri=True,
        )

        try:
            tables = connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table'
                  AND name NOT LIKE 'sqlite_%'
                ORDER BY name
                """
            ).fetchall()

            for (table_name,) in tables:
                count = connection.execute(
                    f'SELECT COUNT(*) FROM "{table_name}"'
                ).fetchone()[0]

                print(f"  {table_name}: {count} records")

        finally:
            connection.close()

    # Inspect the supplied Excel workbook
    excel_path = DATA_DIR / "ridewise_data.xlsx"

    print("\n3. Excel workbook")

    if not excel_path.exists():
        print("  ridewise_data.xlsx was not found.")
        return

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            excel_path,
            read_only=True,
            data_only=True,
        )

        for sheet in workbook.worksheets:
            print(f"\n  Sheet: {sheet.title}")
            print(f"  Rows: {sheet.max_row}")
            print(f"  Columns: {sheet.max_column}")

            headers = next(
                sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                (),
            )

            print("  Headers:")
            print(f"    {list(headers)}")

        workbook.close()

    except Exception as error:
        print(f"  Could not inspect workbook: {error}")


if __name__ == "__main__":
    inspect_files()
